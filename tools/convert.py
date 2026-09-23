#!/usr/bin/env python3
"""
Wandelt die Excel-Medienliste in Dateien für die Suchseite um.

Aufruf (im Hauptordner des Projekts):
    python tools/convert.py                      # liest data/Medienliste.xlsx
    python tools/convert.py pfad/zur/datei.xlsx  # oder eine andere Datei

Erzeugt im Ordner data/:
    medien.json      die bereinigten Daten (für Auswertungen und spätere Erweiterungen)
    medien.js        dieselben Daten, wird von index.html geladen (funktioniert auch per Doppelklick)
    pruefbericht.md  Liste von Auffälligkeiten in der Excel (doppelte Nummern, ungültige ISBN ...)

Benötigt nur:  pip install openpyxl
Die Spalte "Spende" wird bewusst NICHT übernommen.
"""
import datetime as dt
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "Medienliste.xlsx"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "data"

# Spaltenüberschrift in der Excel  ->  Feldname in den Daten.
# Zugeordnet wird über die Überschrift, die Reihenfolge der Spalten ist egal.
# "Spende" fehlt hier absichtlich: Spendernamen sollen nicht veröffentlicht werden.
COLUMNS = {
    "autoren": "author",
    "titel": "title",
    "medien-nr.": "id",
    "signatur": "signature",
    "publikationsdatum": "year",
    "verlag": "publisher",
    "seiten": "pages",
    "isbn": "isbn",
    "standort": "location",
    "anzahl cd": "discs",
    "medienart": "type",
    "sprache": "language",
    "zugang": "added",
    "zugang (datum)": "addedDate",
}
REQUIRED = {"author", "title"}  # Blätter ohne diese Spalten (z. B. "Rückenschilder") werden übersprungen


# ---------------------------------------------------------------- Bereinigung
def text(v):
    """Beliebigen Zellwert in bereinigten Text umwandeln (leer -> None)."""
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:  # NaN
            return None
        if v.is_integer():
            v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def to_int(v):
    if v is None:
        return None
    try:
        return int(float(str(v).replace(",", ".").strip()))
    except ValueError:
        return None


def clean_year(v):
    n = to_int(v)
    if n is not None and 1000 <= n <= 2100:
        return n
    return None  # "k.A." und ähnliches -> keine Angabe


def clean_isbn(v):
    s = text(v)
    if not s:
        return None
    s = re.sub(r"\.0$", "", s)
    s = re.sub(r"[^0-9Xx]", "", s).upper()
    return s or None


def isbn_valid(d):
    if len(d) == 13 and d.isdigit():
        s = sum(int(c) * (1 if i % 2 == 0 else 3) for i, c in enumerate(d[:12]))
        return (10 - s % 10) % 10 == int(d[12])
    if len(d) == 10 and re.fullmatch(r"\d{9}[\dX]", d):
        s = sum((10 - i) * (10 if c == "X" else int(c)) for i, c in enumerate(d))
        return s % 11 == 0
    return False


def iso_date(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return None


# ---------------------------------------------------------------- Einlesen
def read_workbook(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    items, skipped_sheets = [], []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            skipped_sheets.append(ws.title)
            continue
        colmap = {}
        for idx, h in enumerate(header):
            key = COLUMNS.get(str(h).strip().lower()) if h is not None else None
            if key and key not in colmap.values():
                colmap[idx] = key
        if not REQUIRED.issubset(colmap.values()):
            skipped_sheets.append(ws.title)
            continue
        for rownum, row in enumerate(rows, start=2):
            raw = {key: (row[idx] if idx < len(row) else None) for idx, key in colmap.items()}
            if not text(raw.get("title")) and not text(raw.get("author")):
                continue  # leere Zeile
            item = {
                "author": text(raw.get("author")),
                "title": text(raw.get("title")),
                "id": text(raw.get("id")),
                "signature": text(raw.get("signature")),
                "year": clean_year(raw.get("year")),
                "publisher": text(raw.get("publisher")),
                "pages": to_int(raw.get("pages")),
                "isbn": clean_isbn(raw.get("isbn")),
                "location": text(raw.get("location")),
                "discs": text(raw.get("discs")),
                "type": text(raw.get("type")),
                "language": text(raw.get("language")),
            }
            added = text(raw.get("added"))
            item["added"] = None if (added is None or added.lower() == "bestand") else added
            item["addedDate"] = iso_date(raw.get("addedDate"))
            item["_sheet"], item["_row"] = ws.title, rownum
            items.append(item)
    return items, skipped_sheets


# ---------------------------------------------------------------- Prüfbericht
def esc(s, n=70):
    s = (s or "").replace("|", "/")
    return s if len(s) <= n else s[: n - 1] + "…"


def table(rows, head):
    out = ["| " + " | ".join(head) + " |", "|" + "|".join("---" for _ in head) + "|"]
    out += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def fold_block(title, body):
    return f"<details>\n<summary>{title}</summary>\n\n{body}\n\n</details>\n"


def build_report(items, skipped, source_name):
    L = []
    per_sheet = Counter(i["_sheet"] for i in items)
    L.append("# Prüfbericht Medienliste\n")
    L.append(f"Quelle: `{source_name}`, erstellt am {dt.datetime.now().strftime('%d.%m.%Y %H:%M')}\n")
    L.append("Dieser Bericht zeigt Auffälligkeiten in der Excel-Datei. Nichts davon verhindert die Suche, "
             "aber die Punkte lassen sich in der Excel korrigieren. *Zeile* ist die Zeilennummer in Excel.\n")
    L.append("## Übersicht\n")
    L.append(table([(s, n) for s, n in per_sheet.items()] + [("**Gesamt**", f"**{len(items)}**")], ["Blatt", "Einträge"]))
    if skipped:
        L.append(f"\nÜbersprungene Blätter (keine Spalten *Autoren* und *Titel*): {', '.join(skipped)}")

    def ref(i):
        return (i["_sheet"], i["_row"], i["id"] or "–", esc(i["title"]))

    head = ["Blatt", "Zeile", "Medien-Nr.", "Titel"]
    findings = []

    # Medien-Nr.
    missing_id = [i for i in items if not i["id"]]
    by_id = defaultdict(list)
    for i in items:
        if i["id"]:
            by_id[i["id"]].append(i)
    dup_id = {k: v for k, v in by_id.items() if len(v) > 1}
    findings.append(("Medien-Nr. fehlt", len(missing_id), "Ohne Nummer lässt sich ein Medium nicht eindeutig zuordnen.",
                     table([ref(i) for i in missing_id], head) if missing_id else ""))
    rows = [(k, i["_sheet"], i["_row"], esc(i["title"])) for k, v in sorted(dup_id.items()) for i in v]
    findings.append(("Medien-Nr. mehrfach vergeben", len(dup_id),
                     f"{len(dup_id)} Nummern, die bei mehreren Einträgen stehen ({len(rows)} Zeilen).",
                     table(rows, ["Medien-Nr.", "Blatt", "Zeile", "Titel"]) if rows else ""))

    # ISBN
    no_isbn = [i for i in items if not i["isbn"]]
    bad_isbn = [i for i in items if i["isbn"] and not isbn_valid(i["isbn"])]
    by_isbn = defaultdict(list)
    for i in items:
        if i["isbn"]:
            by_isbn[i["isbn"]].append(i)
    dup_isbn = {k: v for k, v in by_isbn.items() if len(v) > 1}
    findings.append(("ISBN fehlt", len(no_isbn), "Diese Medien sind über die ISBN-Suche nicht auffindbar.",
                     table([ref(i) for i in no_isbn], head) if no_isbn else ""))
    findings.append(("ISBN ungültig", len(bad_isbn),
                     "Falsche Länge oder Prüfziffer stimmt nicht (Tippfehler oder Eigencode?). Die Suche findet sie trotzdem.",
                     table([(i["isbn"],) + ref(i) for i in bad_isbn], ["ISBN"] + head) if bad_isbn else ""))
    rows = [(k, i["_sheet"], i["_row"], i["id"] or "–", esc(i["title"])) for k, v in sorted(dup_isbn.items()) for i in v]
    findings.append(("ISBN mehrfach vorhanden", len(dup_isbn),
                     "Kann gewollt sein (mehrere Exemplare), sonst Doppelerfassung.",
                     table(rows, ["ISBN", "Blatt", "Zeile", "Medien-Nr.", "Titel"]) if rows else ""))

    # Medienart / Signatur
    no_type = [i for i in items if not i["type"]]
    by_sheet_notype = Counter(i["_sheet"] for i in no_type)
    body = ""
    if no_type:
        body = table([(s, n) for s, n in by_sheet_notype.items()], ["Blatt", "Einträge ohne Medienart"])
        few = [i for i in no_type if by_sheet_notype[i["_sheet"]] <= 30]
        if few:
            body += "\n\nEinzeln aufgeführt (Blätter mit bis zu 30 Fällen):\n\n" + table([ref(i) for i in few], head)
    findings.append(("Medienart fehlt", len(no_type),
                     "Diese Medien erscheinen nur unter „Ohne Angabe“ im Filter Medienart.", body))

    no_sig = [i for i in items if not i["signature"]]
    findings.append(("Signatur fehlt", len(no_sig), "Ohne Signatur kann das Rückenschild in der Suche nicht angezeigt werden.",
                     table([ref(i) for i in no_sig], head) if no_sig else ""))
    variants = defaultdict(Counter)
    for i in items:
        if i["signature"]:
            variants[i["signature"].lower()][i["signature"]] += 1
    mixed = {k: v for k, v in variants.items() if len(v) > 1}
    findings.append(("Signatur uneinheitlich geschrieben", len(mixed), "Gleiche Signatur mit unterschiedlicher Groß-/Kleinschreibung.",
                     table([(" / ".join(f"{s} ({n}×)" for s, n in v.items()),) for v in mixed.values()], ["Schreibweisen"]) if mixed else ""))

    L.append("\n## Auffälligkeiten\n")
    L.append(table([(t, n) for t, n, _, _ in findings], ["Prüfung", "Anzahl"]) + "\n")
    for t, n, hint, body in findings:
        L.append(f"### {t}: {n}\n")
        L.append(hint + "\n")
        if body:
            L.append(fold_block("Liste anzeigen", body))
    return "\n".join(L), {t: n for t, n, _, _ in findings}


# ---------------------------------------------------------------- Ausgabe
def main():
    if not XLSX.exists():
        sys.exit(f"Datei nicht gefunden: {XLSX}\nBitte die Excel-Liste als data/Medienliste.xlsx ablegen.")
    items, skipped = read_workbook(XLSX)
    if not items:
        sys.exit("Keine Einträge gefunden. Erwartet werden Blätter mit den Spalten 'Autoren' und 'Titel'.")

    report, counts = build_report(items, skipped, XLSX.name)

    public = []
    for i in items:
        public.append({k: v for k, v in i.items() if not k.startswith("_") and v not in (None, "")})
    data = {
        "meta": {
            "generated": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "count": len(public),
            "sheets": dict(Counter(i["_sheet"] for i in items)),
        },
        "items": public,
    }
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "medien.json").write_text(payload, encoding="utf-8")
    (OUT / "medien.js").write_text("window.MEDIEN=" + payload + ";\n", encoding="utf-8")
    (OUT / "pruefbericht.md").write_text(report, encoding="utf-8")

    print(f"{len(public)} Medien aus {len(data['meta']['sheets'])} Blatt/Blättern übernommen -> {OUT}")
    for k, v in counts.items():
        if v:
            print(f"  Hinweis: {k}: {v}")
    print("Details: data/pruefbericht.md")


if __name__ == "__main__":
    main()
